"""Report export: canonical report data rendered to Markdown, DOCX and PDF.

ONE source of truth: the persisted `final_reports.report_markdown` plus the
run's stored trace (citations, sources, claims, metadata). Nothing here
scrapes rendered UI text — the same canonical data the console reads is what
gets exported, so an export can never diverge from the stored report.

Design notes:
  * Markdown is emitted natively: the stored markdown IS the report, so the
    .md export is the canonical body with a metadata front-matter header and
    the source list appended (never a lossy re-render).
  * DOCX/PDF render from a parsed intermediate document so structure
    (headings, tables, lists, citations) is preserved and the typography is
    consistent across both formats.
  * Missing/incomplete reports are handled gracefully: callers get a clear
    ExportError rather than a corrupt file.

Dependencies: python-docx (DOCX) and weasyprint (PDF). Both are imported
lazily inside their renderers so a deployment that only wants Markdown is
not forced to install them, and so the rest of the app imports cleanly when
an optional library is absent.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from html import escape as html_escape
from typing import Any, Dict, List, Optional, Tuple

# Optional deps are resolved at render time (see render_docx/render_pdf).
_SUPPORTED_FORMATS = ("md", "docx", "pdf")

# Machine-readable content types for the HTTP response.
FORMAT_MEDIA_TYPES: Dict[str, str] = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


class ExportError(RuntimeError):
    """Raised when a report cannot be exported (missing/empty/invalid run)."""


@dataclass
class ReportSource:
    marker: Optional[int]
    domain: str
    url: str


@dataclass
class ReportDocument:
    """Format-independent view of a report derived from canonical data."""

    title: str
    query: str
    body_markdown: str
    confidence: Optional[float] = None
    generated_at: str = ""
    run_id: str = ""
    mode: str = ""
    contradiction_count: int = 0
    claim_count: int = 0
    verified_count: int = 0
    sources: List[ReportSource] = field(default_factory=list)

    @property
    def has_body(self) -> bool:
        return bool(self.body_markdown and self.body_markdown.strip())


def is_supported_format(fmt: str) -> bool:
    return str(fmt or "").lower() in _SUPPORTED_FORMATS


def _clean_title(query: str) -> str:
    text = (query or "").strip()
    if not text:
        return "MARS Research Report"
    # Collapse whitespace and cap for use as a filename/title.
    text = re.sub(r"\s+", " ", text)
    return text[:120]


def safe_filename(query: str, run_id: str, fmt: str) -> str:
    """A filesystem-safe download name derived from the query + run id."""
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", (query or "").strip().lower()).strip("-")
    slug = slug[:60] or "report"
    suffix = (run_id or "")[:8]
    stem = f"mars-report-{slug}-{suffix}" if suffix else f"mars-report-{slug}"
    return f"{stem}.{fmt}"


def build_document(trace: Dict[str, Any]) -> ReportDocument:
    """Build a ReportDocument from a stored run trace.

    Raises ExportError when the run is absent or has no completed report, so
    a caller never writes an empty/corrupt file. A report that exists but is
    whitespace-only counts as missing.
    """
    if not trace:
        raise ExportError("Run not found.")
    final = trace.get("final_report") or {}
    body = str(final.get("report_markdown") or "")
    if not body.strip():
        raise ExportError("This run has no completed report to export.")

    confidence = final.get("confidence")
    if confidence is None:
        confidence = trace.get("confidence")

    sources = [
        ReportSource(
            marker=(c.get("marker") if isinstance(c.get("marker"), int) else None),
            domain=str(c.get("domain") or ""),
            url=str(c.get("url") or ""),
        )
        for c in (trace.get("citations") or [])
        if c.get("url")
    ]
    claims = trace.get("claims") or []
    verified = sum(1 for c in claims if c.get("verified") in (1, True))
    contradictions = trace.get("contradictions") or []

    return ReportDocument(
        title=_clean_title(trace.get("query", "")),
        query=str(trace.get("query") or ""),
        body_markdown=body,
        confidence=(float(confidence) if isinstance(confidence, (int, float)) else None),
        generated_at=str(final.get("generated_at") or trace.get("completed_at") or ""),
        run_id=str(trace.get("run_id") or ""),
        mode=str(trace.get("complexity") or ""),
        contradiction_count=len(contradictions),
        claim_count=len(claims),
        verified_count=verified,
        sources=sources,
    )


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def _metadata_lines(doc: ReportDocument) -> List[str]:
    lines = [f"> **Query:** {doc.query}"]
    if doc.confidence is not None:
        lines.append(f"> **Confidence:** {round(doc.confidence * 100)}%")
    if doc.mode:
        lines.append(f"> **Mode:** {doc.mode}")
    lines.append(f"> **Claims:** {doc.claim_count} ({doc.verified_count} verified)")
    if doc.contradiction_count:
        lines.append(f"> **Contradictions:** {doc.contradiction_count}")
    if doc.generated_at:
        lines.append(f"> **Generated:** {doc.generated_at}")
    if doc.run_id:
        lines.append(f"> **Run ID:** {doc.run_id}")
    return lines


def render_markdown(doc: ReportDocument) -> str:
    """Native Markdown: the stored report body verbatim, wrapped with a
    metadata header and an appended source legend when the body lacks one."""
    parts: List[str] = [f"# {doc.title}", ""]
    parts.extend(_metadata_lines(doc))
    parts.extend(["", "---", "", doc.body_markdown.strip()])

    # Preserve citations: only append the source list when the body does not
    # already carry its own legend (avoids duplicating it).
    body_lower = doc.body_markdown.lower()
    already_has_sources = "## sources" in body_lower or "## references" in body_lower
    if doc.sources and not already_has_sources:
        parts.extend(["", "## Sources", ""])
        for src in doc.sources:
            marker = f"[{src.marker}] " if src.marker is not None else ""
            domain = f"{src.domain} — " if src.domain else ""
            parts.append(f"- {marker}{domain}{src.url}")
    return "\n".join(parts).rstrip() + "\n"


# ---------------------------------------------------------------------------
# Markdown → structured blocks (shared by DOCX and PDF)
# ---------------------------------------------------------------------------

@dataclass
class Block:
    kind: str  # heading | paragraph | bullet | numbered | table | code
    text: str = ""
    level: int = 0
    items: List[str] = field(default_factory=list)
    rows: List[List[str]] = field(default_factory=list)


def _strip_inline(text: str) -> str:
    """Reduce inline markdown to plain text for formats that carry their own
    emphasis handling; links become 'label (url)' so citations survive."""
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"(\*\*|__)(.*?)\1", r"\2", text)
    text = re.sub(r"(\*|_)(.*?)\1", r"\2", text)
    text = re.sub(r"^\s*>\s?", "", text)
    return text.strip()


def parse_blocks(markdown: str) -> List[Block]:
    """Lightweight block parser: headings, tables, lists, fenced code and
    paragraphs. Deliberately small — the report is generated by us, so the
    grammar is bounded; unknown lines degrade to paragraphs."""
    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    blocks: List[Block] = []
    para: List[str] = []
    i = 0

    def flush_para() -> None:
        if para:
            text = " ".join(s.strip() for s in para).strip()
            if text:
                blocks.append(Block(kind="paragraph", text=text))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_para()
            code: List[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(Block(kind="code", text="\n".join(code)))
            i += 1
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush_para()
            blocks.append(Block(
                kind="heading", level=len(heading.group(1)), text=heading.group(2).strip(),
            ))
            i += 1
            continue

        # Table: a header row followed by a separator row of dashes/pipes.
        if "|" in stripped and i + 1 < len(lines) and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[i + 1]) and "-" in lines[i + 1]:
            flush_para()
            rows: List[List[str]] = []

            def split_row(row: str) -> List[str]:
                return [c.strip() for c in row.strip().strip("|").split("|")]

            rows.append(split_row(stripped))
            i += 2  # skip header + separator
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(Block(kind="table", rows=rows))
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            flush_para()
            items = [bullet.group(1).strip()]
            i += 1
            while i < len(lines):
                m = re.match(r"^\s*[-*+]\s+(.*)$", lines[i])
                if not m:
                    break
                items.append(m.group(1).strip())
                i += 1
            blocks.append(Block(kind="bullet", items=items))
            continue

        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            flush_para()
            items = [numbered.group(1).strip()]
            i += 1
            while i < len(lines):
                m = re.match(r"^\s*\d+[.)]\s+(.*)$", lines[i])
                if not m:
                    break
                items.append(m.group(1).strip())
                i += 1
            blocks.append(Block(kind="numbered", items=items))
            continue

        para.append(stripped)
        i += 1

    flush_para()
    return blocks


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def render_docx(doc: ReportDocument) -> bytes:
    """Professional DOCX with styled headings, tables and a source list."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except Exception as exc:  # pragma: no cover - optional dependency
        raise ExportError(
            "DOCX export is unavailable: the 'python-docx' package is not installed."
        ) from exc

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = document.add_heading(doc.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = document.add_paragraph()
    meta_run = meta.add_run(
        " | ".join(
            filter(None, [
                f"Confidence: {round(doc.confidence * 100)}%" if doc.confidence is not None else "",
                f"Mode: {doc.mode}" if doc.mode else "",
                f"Claims: {doc.claim_count} ({doc.verified_count} verified)",
                f"Generated: {doc.generated_at}" if doc.generated_at else "",
            ])
        )
    )
    meta_run.italic = True
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    document.add_paragraph()

    def add_runs(paragraph, text: str) -> None:
        """Render **bold** spans instead of dropping them."""
        for idx, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
            if not chunk:
                continue
            if chunk.startswith("**") and chunk.endswith("**"):
                run = paragraph.add_run(chunk[2:-2])
                run.bold = True
            else:
                paragraph.add_run(_strip_inline(chunk))

    for block in parse_blocks(doc.body_markdown):
        if block.kind == "heading":
            level = min(max(block.level, 1), 4)
            document.add_heading(_strip_inline(block.text), level=level)
        elif block.kind == "paragraph":
            add_runs(document.add_paragraph(), block.text)
        elif block.kind == "bullet":
            for item in block.items:
                add_runs(document.add_paragraph(style="List Bullet"), item)
        elif block.kind == "numbered":
            for item in block.items:
                add_runs(document.add_paragraph(style="List Number"), item)
        elif block.kind == "code":
            para = document.add_paragraph()
            run = para.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif block.kind == "table" and block.rows:
            cols = max(len(r) for r in block.rows)
            table = document.add_table(rows=0, cols=cols)
            table.style = "Light Grid Accent 1"
            for row in block.rows:
                cells = table.add_row().cells
                for j in range(cols):
                    cells[j].text = _strip_inline(row[j]) if j < len(row) else ""

    if doc.sources:
        document.add_heading("Sources", level=1)
        for src in doc.sources:
            marker = f"[{src.marker}] " if src.marker is not None else ""
            domain = f"{src.domain} — " if src.domain else ""
            document.add_paragraph(f"{marker}{domain}{src.url}", style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_PDF_CSS = """
@page { size: A4; margin: 20mm 18mm; }
* { box-sizing: border-box; }
body { font-family: "DejaVu Sans", "Helvetica", sans-serif; color: #1a1a1a;
       font-size: 10.5pt; line-height: 1.55; }
h1 { font-size: 20pt; margin: 0 0 6pt; color: #111; }
h2 { font-size: 15pt; margin: 18pt 0 6pt; color: #b8532a;
     border-bottom: 1px solid #e6d8cf; padding-bottom: 3pt; }
h3 { font-size: 12.5pt; margin: 14pt 0 5pt; color: #222; }
h4 { font-size: 11pt; margin: 12pt 0 4pt; color: #333; }
p { margin: 0 0 8pt; }
ul, ol { margin: 0 0 9pt; padding-left: 18pt; }
li { margin: 2pt 0; }
.meta { color: #666; font-size: 8.5pt; margin: 0 0 14pt;
        border-bottom: 2px solid #b8532a; padding-bottom: 8pt; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0 12pt; }
th, td { border: 1px solid #ccc; padding: 5pt 7pt; text-align: left;
         font-size: 9.5pt; }
th { background: #f4ede8; font-weight: 600; }
pre { background: #f5f5f5; border: 1px solid #e0e0e0; border-radius: 4px;
      padding: 8pt; font-family: "DejaVu Sans Mono", monospace; font-size: 8.5pt;
      white-space: pre-wrap; }
a { color: #b8532a; text-decoration: none; word-break: break-all; }
.sources li { font-size: 9pt; }
"""


def _html_for_document(doc: ReportDocument) -> str:
    out: List[str] = ["<!DOCTYPE html><html><head><meta charset='utf-8'>"]
    out.append(f"<style>{_PDF_CSS}</style></head><body>")
    out.append(f"<h1>{html_escape(doc.title)}</h1>")
    meta_bits = [
        f"Confidence: {round(doc.confidence * 100)}%" if doc.confidence is not None else "",
        f"Mode: {html_escape(doc.mode)}" if doc.mode else "",
        f"Claims: {doc.claim_count} ({doc.verified_count} verified)",
        f"Generated: {html_escape(doc.generated_at)}" if doc.generated_at else "",
    ]
    out.append(f"<div class='meta'>{' &middot; '.join(b for b in meta_bits if b)}</div>")

    for block in parse_blocks(doc.body_markdown):
        if block.kind == "heading":
            level = min(max(block.level, 1), 4)
            out.append(f"<h{level}>{html_escape(_strip_inline(block.text))}</h{level}>")
        elif block.kind == "paragraph":
            out.append(f"<p>{html_escape(_strip_inline(block.text))}</p>")
        elif block.kind == "bullet":
            items = "".join(f"<li>{html_escape(_strip_inline(x))}</li>" for x in block.items)
            out.append(f"<ul>{items}</ul>")
        elif block.kind == "numbered":
            items = "".join(f"<li>{html_escape(_strip_inline(x))}</li>" for x in block.items)
            out.append(f"<ol>{items}</ol>")
        elif block.kind == "code":
            out.append(f"<pre>{html_escape(block.text)}</pre>")
        elif block.kind == "table" and block.rows:
            rows_html: List[str] = []
            for r_idx, row in enumerate(block.rows):
                cell_tag = "th" if r_idx == 0 else "td"
                cells = "".join(
                    f"<{cell_tag}>{html_escape(_strip_inline(c))}</{cell_tag}>" for c in row
                )
                rows_html.append(f"<tr>{cells}</tr>")
            out.append(f"<table>{''.join(rows_html)}</table>")
    if doc.sources:
        out.append("<h2>Sources</h2><ul class='sources'>")
        for src in doc.sources:
            marker = f"[{src.marker}] " if src.marker is not None else ""
            label = f"{marker}{html_escape(src.domain)} — " if src.domain else marker
            safe_url = html_escape(src.url)
            out.append(f"<li>{label}<a href='{safe_url}'>{safe_url}</a></li>")
        out.append("</ul>")
    out.append("</body></html>")
    return "".join(out)


def render_pdf(doc: ReportDocument) -> bytes:
    """Professional PDF via WeasyPrint from the structured document."""
    try:
        from weasyprint import HTML
    except Exception as exc:  # pragma: no cover - optional dependency
        raise ExportError(
            "PDF export is unavailable: the 'weasyprint' package is not installed."
        ) from exc
    try:
        return HTML(string=_html_for_document(doc)).write_pdf()
    except Exception as exc:  # pragma: no cover - render-time failure
        raise ExportError(f"PDF rendering failed: {type(exc).__name__}") from exc


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

def render_report(trace: Dict[str, Any], fmt: str) -> Tuple[bytes, str, str]:
    """Render a stored trace to (bytes, media_type, filename).

    Raises ExportError for an unsupported format or a missing/empty report.
    """
    fmt = str(fmt or "").lower()
    if not is_supported_format(fmt):
        raise ExportError(f"Unsupported export format: {fmt or '(none)'}")
    doc = build_document(trace)
    if fmt == "md":
        return render_markdown(doc).encode("utf-8"), FORMAT_MEDIA_TYPES["md"], safe_filename(doc.query, doc.run_id, "md")
    if fmt == "docx":
        return render_docx(doc), FORMAT_MEDIA_TYPES["docx"], safe_filename(doc.query, doc.run_id, "docx")
    return render_pdf(doc), FORMAT_MEDIA_TYPES["pdf"], safe_filename(doc.query, doc.run_id, "pdf")
